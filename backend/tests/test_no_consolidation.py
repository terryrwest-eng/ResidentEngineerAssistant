"""Resource rows are listed as entered, never merged.

WHY THIS EXISTS: the Word report used to group rows by (resource, hours,
company) and sum the quantities, so a shift with eight named labourers printed
as one line reading "LL-03- Laborers - QTY 8", and two traffic control trucks
printed as "QTY 2". The names were entered by hand and then thrown away by the
renderer, which reads exactly like the app losing your work.

The merging lived in three places that had to agree — the Word activity list,
the "Consolidated Resources" grid at the end of the document, and the PMWeb
Notes HTML. The grid is gone and the other two now share one line builder, so
they cannot drift apart again.
"""
import io
import sys

sys.path.insert(0, "backend")

from app.services.word import generate_word_document, generate_notes_html  # noqa: E402

import docx  # noqa: E402

results = []


def check(name, cond, detail=""):
    results.append(cond)
    print(f"{'PASS' if cond else 'FAIL'}  {name}" + (f"  — {detail}" if detail else ""))


REPORT = {
    "general": {
        "project_name": "Morena Conveyance Northern",
        "report_date": "2026-08-05",
        "inspector_name": "Terry West",
        "start_time": "11:30 PM",
        "end_time": "3:45 AM",
    },
    "activities": [
        {
            "work_area": "Genesee, Towne Centre, Nobel, Executive",
            "summary": "• Striping operations.",
            "manpower": [
                {"trade": "LL-02- Foreman", "name": "Stephen Strandberg",
                 "qty": 1, "hours": 4.25, "company": "Payco"},
                {"trade": "LL-03- Laborers", "name": "Dan Griffin",
                 "qty": 1, "hours": 4.25, "company": "Payco"},
                {"trade": "LL-03- Laborers", "name": "Juan Medina",
                 "qty": 1, "hours": 4.25, "company": "Payco"},
            ],
            "equipment": [
                {"name": "LE-161- Traffic Control Truck", "qty": 1,
                 "hours": 4.25, "company": "Payco"},
                {"name": "LE-161- Traffic Control Truck", "qty": 1,
                 "hours": 4.25, "company": "Payco"},
                {"name": "LE-170- Airless Paint Striper", "qty": 1,
                 "hours": 4.25, "company": "Payco"},
            ],
        }
    ],
}


def doc_lines():
    d = docx.Document(generate_word_document(REPORT))
    return [p.text.strip() for p in d.paragraphs if p.text.strip()]


lines = doc_lines()
body = "\n".join(lines)

# --- rows are not merged ---------------------------------------------------
check("both labourers get their own line",
      sum(1 for ln in lines if ln.startswith("LL-03- Laborers")) == 2,
      f"found {sum(1 for ln in lines if ln.startswith('LL-03- Laborers'))}")

check("both traffic control trucks get their own line",
      sum(1 for ln in lines if ln.startswith("LE-161- Traffic Control Truck")) == 2,
      f"found {sum(1 for ln in lines if ln.startswith('LE-161- Traffic Control Truck'))}")

check("no row is summed into a quantity it was not given",
      "QTY 2" not in body and "QTY 3" not in body)

# --- names survive ---------------------------------------------------------
for who in ("Stephen Strandberg", "Dan Griffin", "Juan Medina"):
    check(f"name kept: {who}", who in body)

# --- equipment keeps the name it was given ---------------------------------
check("paint striper is not renamed to a person",
      "LE-170- Airless Paint Striper" in body and "LL-09- PE" not in body)

# --- the consolidated grid is gone -----------------------------------------
d = docx.Document(generate_word_document(REPORT))
check("no Consolidated Resources heading", "Consolidated Resources" not in body)
check("document has no tables at all", len(d.tables) == 0,
      f"{len(d.tables)} table(s)")

# --- the PMWeb Notes HTML matches the document -----------------------------
html = generate_notes_html(REPORT)
check("notes HTML lists both labourers",
      html.count("LL-03- Laborers") == 2, f"found {html.count('LL-03- Laborers')}")
check("notes HTML lists both trucks",
      html.count("LE-161- Traffic Control Truck") == 2,
      f"found {html.count('LE-161- Traffic Control Truck')}")
check("notes HTML keeps the names", "Dan Griffin" in html and "Juan Medina" in html)

# --- "EA" only when a row really covers more than one ----------------------
multi = dict(REPORT)
multi = {
    **REPORT,
    "activities": [
        {
            **REPORT["activities"][0],
            "manpower": [{"trade": "LL-03- Laborers", "name": "Crew",
                          "qty": 3, "hours": 4.25, "company": "Payco"}],
            "equipment": [],
        }
    ],
}
multi_body = "\n".join(
    p.text.strip() for p in docx.Document(generate_word_document(multi)).paragraphs
)
# Not asserting the rendered number itself — _format_number owns that.
check("a qty>1 row still reads HRS EA",
      "QTY 3" in multi_body and "HRS EA" in multi_body,
      multi_body.split("Manpower:")[-1].strip()[:60])
check("a qty=1 row reads plain HRS, not HRS EA",
      "HRS -" in body and "HRS EA" not in body)

print(f"\n{sum(results)}/{len(results)} passed")
sys.exit(0 if all(results) else 1)
